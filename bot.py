import os
import random
import sqlite3
import threading
from flask import Flask
from openpyxl import load_workbook
import telebot
from telebot import types
import google.generativeai as genai

# --- BOT VE ÇEVRESEL AYARLAR ---
TOKEN = os.environ.get("BOT_TOKEN", "8899348135:AAH2G6dvKUSaSXuNI0p1PT3_JDCMgjRXaGo")
GEMINI_KEY = os.environ.get("GEMINI_API_KEY", "")

bot = telebot.TeleBot(TOKEN)

ALLOWED_USERS = [8022316805, 0]

def is_authorized(user_id):
    return user_id in ALLOWED_USERS

if GEMINI_KEY:
    try:
        genai.configure(api_key=GEMINI_KEY)
        ai_model = genai.GenerativeModel("gemini-1.5-flash")
    except Exception as e:
        ai_model = None
else:
    ai_model = None

DB_FILE = "english_bot.db"
EXCEL_FILE = "translated.xlsx"

active_placement_tests = {}
ai_chat_users = set()

# --- FLASK WEB SUNUCUSU ---
app = Flask(__name__)

@app.route('/')
def home():
    return "English Coach AI Platform Aktif!"

@app.route('/health')
def health():
    return "OK", 200

def run_flask():
    port = int(os.environ.get("PORT", 8080))
    app.run(host="0.0.0.0", port=port)

# --- SORULAR VE DERSLER (Alt tireler Markdown bozmaması için nokta yapıldı) ---

PLACEMENT_QUESTIONS = [
    {"q": "She ....... from Spain.", "options": ["is", "are", "am", "be"], "a": "is", "lvl": "A1"},
    {"q": "I ....... a movie yesterday night.", "options": ["watched", "watch", "watching", "have watched"], "a": "watched", "lvl": "A1"},
    {"q": "There aren't ....... apples left in the basket.", "options": ["any", "some", "much", "an"], "a": "any", "lvl": "A2"},
    {"q": "He was cooking dinner while she ....... a book.", "options": ["was reading", "reads", "has read", "is reading"], "a": "was reading", "lvl": "A2"},
    {"q": "If it rains tomorrow, we ....... the match.", "options": ["will cancel", "canceled", "cancel", "would cancel"], "a": "will cancel", "lvl": "B1"},
    {"q": "I have lived in this city ....... 2018.", "options": ["since", "for", "in", "during"], "a": "since", "lvl": "B1"},
    {"q": "By the time the police arrived, the thieves ....... .", "options": ["had escaped", "escaped", "have escaped", "escape"], "a": "had escaped", "lvl": "B2"},
    {"q": "If I ....... you, I would take that offer immediately.", "options": ["were", "am", "will be", "had been"], "a": "were", "lvl": "B2"},
    {"q": "Seldom ....... such an extraordinary performance.", "options": ["have I seen", "I saw", "I have seen", "saw I"], "a": "have I seen", "lvl": "C1"},
    {"q": "The manager insisted that the report ....... submitted by noon.", "options": ["be", "is", "was", "will be"], "a": "be", "lvl": "C1"}
]

CATEGORY_DATA = {
    "travel": {
        "title": "✈️ Seyahat & Tatil",
        "words": [("passport", "pasaport"), ("boarding pass", "biniş kartı"), ("customs", "gümrük"), ("baggage claim", "bagaj alım yeri"), ("delay", "gecikme"), ("reservation", "rezervasyon")],
        "sentences": [("Where is the baggage claim area?", "Bagaj alım alanı nerede?"), ("Here is my passport and boarding pass.", "Buyrun pasaportum ve biniş kartım.")]
    },
    "business": {
        "title": "💼 İş Dünyası",
        "words": [("deadline", "son teslim tarihi"), ("meeting", "toplantı"), ("negotiate", "müzakere etmek"), ("salary", "maaş")],
        "sentences": [("We need to meet the deadline tomorrow.", "Yarınki son teslim tarihine yetişmemiz gerekiyor.")]
    },
    "food": {
        "title": "🍽 Restoran & Yemek",
        "words": [("menu", "menü"), ("bill", "hesap"), ("waiter", "garson"), ("delicious", "lezzetli")],
        "sentences": [("Could we have the bill, please?", "Hesabı alabilir miyiz lütfen?")]
    },
    "daily": {
        "title": "🏫 Günlük Yaşam",
        "words": [("neighborhood", "mahalle"), ("groceries", "market alışverişi"), ("commute", "işe/okula gidip gelmek")],
        "sentences": [("What time does the bus leave?", "Otobüs saat kaçta kalkıyor?")]
    }
}

GRAMMAR_LESSONS = {
    "A1": "📘 A1: To Be (am / is / are)\n\n• I -> am\n• He / She / It -> is\n• You / We / They -> are\n\nÖrnek: She is a doctor.",
    "A2": "📗 A2: Past Simple (Geçmiş Zaman)\n\n• Düzenli: -ed (watched)\n• Düzensiz: (went, saw)\n\nÖrnek: I went home.",
    "B1": "📙 B1: Present Perfect Tense\n\n• have/has + V3\n\nÖrnek: I have visited London.",
    "B2": "📕 B2: Conditionals\n\n• Type 2: If I were you, I would study."
}

# --- VERİTABANI ---

def init_database():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS words (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            english TEXT UNIQUE,
            turkish TEXT
        )
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS user_profiles (
            user_id INTEGER PRIMARY KEY,
            level TEXT DEFAULT 'A1',
            correct_count INTEGER DEFAULT 0,
            wrong_count INTEGER DEFAULT 0
        )
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS unknown_words (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            english TEXT,
            turkish TEXT,
            UNIQUE(user_id, english)
        )
    """)
    cursor.execute("SELECT COUNT(*) FROM words")
    count = cursor.fetchone()[0]
    if count == 0 and os.path.exists(EXCEL_FILE):
        try:
            wb = load_workbook(EXCEL_FILE, read_only=True)
            sheet = wb.active
            for row in list(sheet.iter_rows(values_only=True))[1:]:
                if row and len(row) >= 2:
                    eng = str(row[0]).strip() if row[0] is not None else ""
                    tr = str(row[1]).strip() if row[1] is not None else ""
                    if eng and tr and eng.lower() != 'the oxford':
                        cursor.execute("INSERT OR IGNORE INTO words (english, turkish) VALUES (?, ?)", (eng, tr))
            conn.commit()
        except Exception as e:
            print("Excel hatasi:", e)
    conn.close()

def get_user_profile(user_id):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("INSERT OR IGNORE INTO user_profiles (user_id, level, correct_count, wrong_count) VALUES (?, 'A1', 0, 0)", (user_id,))
    conn.commit()
    cursor.execute("SELECT level, correct_count, wrong_count FROM user_profiles WHERE user_id = ?", (user_id,))
    res = cursor.fetchone()
    conn.close()
    return res

def set_user_level(user_id, level):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("UPDATE user_profiles SET level = ? WHERE user_id = ?", (level, user_id))
    conn.commit()
    conn.close()

def update_score(user_id, is_correct):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    if is_correct:
        cursor.execute("UPDATE user_profiles SET correct_count = correct_count + 1 WHERE user_id = ?", (user_id,))
    else:
        cursor.execute("UPDATE user_profiles SET wrong_count = wrong_count + 1 WHERE user_id = ?", (user_id,))
    conn.commit()
    conn.close()

def add_unknown_word(user_id, eng, tr):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("INSERT OR IGNORE INTO unknown_words (user_id, english, turkish) VALUES (?, ?, ?)", (user_id, eng, tr))
    conn.commit()
    conn.close()

def remove_unknown_word(user_id, eng):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM unknown_words WHERE user_id = ? AND english = ?", (user_id, eng))
    conn.commit()
    conn.close()

def get_unknown_words_count(user_id):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM unknown_words WHERE user_id = ?", (user_id,))
    count = cursor.fetchone()[0]
    conn.close()
    return count

def get_random_unknown_quiz(user_id):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("SELECT english, turkish FROM unknown_words WHERE user_id = ? ORDER BY RANDOM() LIMIT 1", (user_id,))
    target = cursor.fetchone()
    if not target:
        conn.close()
        return None, None, []
    cursor.execute("SELECT turkish FROM words WHERE english != ? ORDER BY RANDOM() LIMIT 3", (target[0],))
    distractors = [r[0] for r in cursor.fetchall()]
    conn.close()
    options = distractors + [target[1]]
    random.shuffle(options)
    return target[0], target[1], options

def get_general_quiz():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("SELECT english, turkish FROM words ORDER BY RANDOM() LIMIT 4")
    rows = cursor.fetchall()
    conn.close()
    if len(rows) < 4:
        return None, None, []
    target_eng, target_tr = rows[0]
    options = [r[1] for r in rows]
    random.shuffle(options)
    return target_eng, target_tr, options

# --- ANA KLAVYE ---

def get_main_keyboard():
    keyboard = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    keyboard.add(types.KeyboardButton("🤖 AI Koç"))
    keyboard.add(
        types.KeyboardButton("📚 Kelime Çöz"),
        types.KeyboardButton("🧠 Zayıf Kelimelerim")
    )
    keyboard.add(
        types.KeyboardButton("🗂 Kategorik Kelimeler"),
        types.KeyboardButton("📖 Gramer Dersleri")
    )
    keyboard.add(
        types.KeyboardButton("🎓 Seviye Belirleme Sınavı"),
        types.KeyboardButton("📊 Profilim & Skor")
    )
    return keyboard

# --- SEVİYE SINAVI ÇALIŞTIRICI ---

def run_placement_question(chat_id, user_id):
    try:
        state = active_placement_tests.get(user_id)
        if not state:
            state = {"q_idx": 0, "score": 0}
            active_placement_tests[user_id] = state

        idx = state["q_idx"]
        if idx >= len(PLACEMENT_QUESTIONS):
            score = state["score"]
            if score <= 2: final_lvl = "A1"
            elif score <= 4: final_lvl = "A2"
            elif score <= 7: final_lvl = "B1"
            elif score <= 9: final_lvl = "B2"
            else: final_lvl = "C1"

            set_user_level(user_id, final_lvl)
            active_placement_tests.pop(user_id, None)
            bot.send_message(
                chat_id,
                f"🏁 Seviye Belirleme Sınavı Tamamlandı!\n\n"
                f"✅ Doğru Sayın: {score} / 10\n"
                f"🎯 Belirlenen Seviyen: {final_lvl}\n\n"
                f"Profilin güncellendi!",
                reply_markup=get_main_keyboard()
            )
            return

        q_data = PLACEMENT_QUESTIONS[idx]
        markup = types.InlineKeyboardMarkup(row_width=2)
        for i, opt in enumerate(q_data["options"]):
            markup.add(types.InlineKeyboardButton(opt, callback_data=f"PT|{i}"))

        # parse_mode kaldırıldı, alt tire ve işaretler Telegram'ı patlatmayacak
        bot.send_message(
            chat_id,
            f"📝 Soru {idx + 1}/10 ({q_data['lvl']}):\n\n{q_data['q']}",
            reply_markup=markup
        )
    except Exception as e:
        bot.send_message(chat_id, f"Sınav başlatılırken hata oluştu: {e}")

# --- MESAJ YÖNLENDİRİCİ ---

@bot.message_handler(commands=['start'])
def handle_start(message):
    uid = message.from_user.id
    if not is_authorized(uid):
        bot.send_message(message.chat.id, "⛔ Erişim Engellendi: Bu bot sadece yetkili kullanıcılara özeldir.")
        return

    get_user_profile(uid)
    ai_chat_users.discard(uid)
    welcome_text = (
        f"Selam {message.from_user.first_name}! 🇬🇧\n\n"
        "AI Destekli İngilizce Koçuna hoş geldin! Aşağıdaki butonları kullanarak hemen başlayabilirsin."
    )
    bot.send_message(message.chat.id, welcome_text, reply_markup=get_main_keyboard())

@bot.message_handler(func=lambda msg: True)
def handle_text_menu(message):
    uid = message.from_user.id
    if not is_authorized(uid):
        return

    raw_text = (message.text or "")
    cleaned = raw_text.replace("\n", " ").strip().lower()

    # 1. SEVİYE BELİRLEME SINAVI
    if "seviye" in cleaned or "sınav" in cleaned or "sinav" in cleaned:
        ai_chat_users.discard(uid)
        active_placement_tests[uid] = {"q_idx": 0, "score": 0}
        run_placement_question(message.chat.id, uid)
        return

    # 2. AI KOÇ
    if "ai koç" in cleaned or "ai koc" in cleaned:
        ai_chat_users.add(uid)
        bot.send_message(
            message.chat.id,
            "🤖 AI İngilizce Koçun Dinliyor!\n\n"
            "Bana istediğin İngilizce cümleyi veya soruyu yazabilirsin.\n"
            "✍️ Mesajını buraya yaz:"
        )
        return

    # 3. KELİME ÇÖZ
    if "kelime" in cleaned and ("çöz" in cleaned or "coz" in cleaned):
        ai_chat_users.discard(uid)
        send_quiz_question(message.chat.id, is_unknown_mode=False)
        return

    # 4. ZAYIF KELİMELERİM
    if "zayıf" in cleaned or "zayif" in cleaned:
        ai_chat_users.discard(uid)
        count = get_unknown_words_count(uid)
        if count == 0:
            bot.send_message(message.chat.id, "🎉 Harika! Zayıf kelimeler listende henüz kelime yok.")
        else:
            send_quiz_question(message.chat.id, is_unknown_mode=True)
        return

    # 5. KATEGORİK KELİMELER
    if "kategori" in cleaned:
        ai_chat_users.discard(uid)
        markup = types.InlineKeyboardMarkup(row_width=2)
        for cat_key, cat_val in CATEGORY_DATA.items():
            markup.add(
                types.InlineKeyboardButton(f"{cat_val['title']} (Kelimeler)", callback_data=f"cw:{cat_key}"),
                types.InlineKeyboardButton("🗣 Cümleler", callback_data=f"cs:{cat_key}")
            )
        bot.send_message(message.chat.id, "📌 Çalışmak istediğin alanı seç:", reply_markup=markup)
        return

    # 6. GRAMER DERSLERİ
    if "gramer" in cleaned:
        ai_chat_users.discard(uid)
        markup = types.InlineKeyboardMarkup(row_width=2)
        markup.add(
            types.InlineKeyboardButton("A1: To Be", callback_data="gm:A1"),
            types.InlineKeyboardButton("A2: Past Simple", callback_data="gm:A2"),
            types.InlineKeyboardButton("B1: Present Perfect", callback_data="gm:B1"),
            types.InlineKeyboardButton("B2: Conditionals", callback_data="gm:B2")
        )
        bot.send_message(message.chat.id, "📘 Gramer Seviyesini Seç:", reply_markup=markup)
        return

    # 7. PROFİLİM & SKOR
    if "profil" in cleaned or "skor" in cleaned:
        ai_chat_users.discard(uid)
        prof = get_user_profile(uid)
        lvl, c, w = prof[0], prof[1], prof[2]
        un_count = get_unknown_words_count(uid)
        total = c + w
        rate = f"%{int((c / total) * 100)}" if total > 0 else "%0"
        msg = (
            f"👤 Öğrenci Profili\n\n"
            f"🎯 Seviyen: {lvl}\n"
            f"✅ Doğru: {c}\n"
            f"❌ Yanlış: {w}\n"
            f"📈 Başarı: {rate}\n"
            f"🧠 Tekrar Bekleyen Kelimeler: {un_count} adet"
        )
        bot.send_message(message.chat.id, msg)
        return

    # 8. AI SOHBET
    if uid in ai_chat_users:
        if not ai_model:
            bot.send_message(message.chat.id, "⚠️ AI anahtarı tanımlanmamış.")
            return

        bot.send_chat_action(message.chat.id, 'typing')
        prompt = (
            "Sen samimi ve uzman bir İngilizce öğretmenisin. "
            "Kullanıcının yazdığı İngilizce cümleyi veya soruyu analiz et. "
            "Hata varsa doğrusunu göster ve nedenini Türkçe olarak açıkla. "
            f"Kullanıcı mesajı: {raw_text}"
        )
        try:
            response = ai_model.generate_content(prompt)
            bot.reply_to(message, response.text)
        except Exception as e:
            bot.reply_to(message, f"Hata: {e}")

# --- KELİME TESTİ ---

def send_quiz_question(chat_id, is_unknown_mode=False):
    if is_unknown_mode:
        eng, tr, options = get_random_unknown_quiz(chat_id)
        mode = "U"
    else:
        eng, tr, options = get_general_quiz()
        mode = "G"

    if not eng:
        bot.send_message(chat_id, "Kelime bulunamadı, lütfen tekrar deneyin.")
        return

    markup = types.InlineKeyboardMarkup(row_width=1)
    for opt in options:
        status = "C" if opt == tr else "W"
        cb = f"Q|{mode}|{status}|{eng[:10]}|{tr[:10]}"
        markup.add(types.InlineKeyboardButton(opt, callback_data=cb))
    
    markup.add(types.InlineKeyboardButton("🤷‍♂️ Bilmiyorum", callback_data=f"Q|{mode}|IDK|{eng[:10]}|{tr[:10]}"))
    prefix = "🧠 [Zayıf Kelime]\n" if is_unknown_mode else "🎯 "
    bot.send_message(chat_id, f"{prefix}{eng} kelimesinin Türkçe karşılığı nedir?", reply_markup=markup)

# --- BUTON CEVAPLARI ---

@bot.callback_query_handler(func=lambda call: True)
def handle_callbacks(call):
    uid = call.from_user.id
    if not is_authorized(uid):
        bot.answer_callback_query(call.id, "⛔ Yetkisiz erişim!", show_alert=True)
        return

    bot.answer_callback_query(call.id)
    data = call.data
    cid = call.message.chat.id
    mid = call.message.message_id

    # Sınav Şıkları
    if data.startswith("PT|"):
        opt_idx = int(data.split("|")[1])
        state = active_placement_tests.get(uid)
        if state:
            idx = state["q_idx"]
            if idx < len(PLACEMENT_QUESTIONS):
                q_data = PLACEMENT_QUESTIONS[idx]
                if q_data["options"][opt_idx] == q_data["a"]:
                    state["score"] += 1
                state["q_idx"] += 1
                try:
                    bot.delete_message(cid, mid)
                except:
                    pass
                run_placement_question(cid, uid)

    # Kelime Testi Şıkları
    elif data.startswith("Q|"):
        _, mode, status, eng, tr = data.split("|", 4)
        if status == "C":
            update_score(uid, True)
            extra = "\n✨ Bu kelimeyi öğrendin, listeden kaldırıldı!" if mode == "U" else ""
            if mode == "U":
                remove_unknown_word(uid, eng)
            bot.edit_message_text(f"✅ Doğru!\n\n🇬🇧 {eng} = 🇹🇷 {tr}{extra}", chat_id=cid, message_id=mid)
        elif status == "W":
            update_score(uid, False)
            add_unknown_word(uid, eng, tr)
            bot.edit_message_text(f"❌ Yanlış!\n\n🇬🇧 {eng} kelimesinin doğrusu: 🇹🇷 {tr}\n\n📌 Zayıf kelimeler listene eklendi.", chat_id=cid, message_id=mid)
        elif status == "IDK":
            add_unknown_word(uid, eng, tr)
            bot.edit_message_text(f"💡 Doğrusu:\n\n🇬🇧 {eng} = 🇹🇷 {tr}\n\n📌 Zayıf kelimeler listene eklendi.", chat_id=cid, message_id=mid)

    elif data.startswith("cw:"):
        k = data.split(":")[1]
        cat = CATEGORY_DATA.get(k)
        if cat:
            txt = f"{cat['title']} - Kelimeler:\n\n" + "\n".join([f"• {e}: {t}" for e, t in cat['words']])
            bot.send_message(cid, txt)

    elif data.startswith("cs:"):
        k = data.split(":")[1]
        cat = CATEGORY_DATA.get(k)
        if cat:
            txt = f"{cat['title']} - Kalıp Cümleler:\n\n" + "\n\n".join([f"💬 {e}\n↳ {t}" for e, t in cat['sentences']])
            bot.send_message(cid, txt)

    elif data.startswith("gm:"):
        lvl = data.split(":")[1]
        lesson = GRAMMAR_LESSONS.get(lvl)
        if lesson:
            bot.send_message(cid, lesson)

# --- BAŞLATMA ---

if __name__ == "__main__":
    init_database()
    t = threading.Thread(target=run_flask)
    t.daemon = True
    t.start()
    print("Bot basariyla baslatildi...")
    bot.infinity_polling()
